from io import BytesIO
from types import SimpleNamespace
from typing import cast

import pytest
import pymupdf
from docx import Document

from app.api.admin.documents import SUPPORTED_UPLOAD_EXTENSIONS
from app.api.admin.ingestion import _resolve_source_type
from app.api.dependencies import get_knowledge_ingestion_pipeline
from app.config.settings import load_settings
from app.core.exceptions import ApplicationError
from app.infrastructure.providers.parsing.docling_provider import (
    DoclingParsingProvider,
)
from app.infrastructure.providers.parsing.pymupdf_provider import (
    PyMuPDFParsingProvider,
)
from app.infrastructure.providers.parsing.python_docx_provider import (
    PythonDocxParsingProvider,
)
from app.knowledge_engine.contracts.embeddings import EmbeddingsContract
from app.knowledge_engine.contracts.vector_store import VectorStoreContract
from app.knowledge_engine.ingestion.parsers.csv_parser import CsvDocumentParser
from app.knowledge_engine.ingestion.parsers.structured_document_parser import (
    StructuredDocumentParser,
)
from app.knowledge_engine.ingestion.parsers.text_parser import TextDocumentParser
from app.knowledge_engine.ingestion.source_loaders.file_loader import FileSourceLoader
from app.knowledge_engine.shared.models import (
    KnowledgeIngestionPipelineRequest,
    ParsedDocument,
    RawSource,
)


def _pipeline_for(source_type: str):
    settings = SimpleNamespace(
        providers=SimpleNamespace(
            embeddings="nomic",
            parser="native",
        ),
        storage=SimpleNamespace(),
        openrouter=SimpleNamespace(),
        vector_store_table_name="chunks",
        vector_store_dimension=3,
    )
    request = SimpleNamespace(
        app=SimpleNamespace(state=SimpleNamespace(settings=settings))
    )
    return get_knowledge_ingestion_pipeline(
        source_type=source_type,
        request=request,
        session=object(),
    )


def test_docx_uses_python_docx_parser():
    pipeline = _pipeline_for("docx")

    assert isinstance(pipeline.source_loader, FileSourceLoader)
    assert isinstance(pipeline.parser, StructuredDocumentParser)
    assert isinstance(
        pipeline.parser.parsing_contract,
        PythonDocxParsingProvider,
    )


def test_docling_provider_routes_pdf_and_docx_when_configured():
    for source_type in ("pdf", "docx"):
        pipeline = _pipeline_for_provider(source_type, "docling")

        assert isinstance(pipeline.parser, StructuredDocumentParser)
        assert isinstance(pipeline.parser.parsing_contract, DoclingParsingProvider)


def test_docling_enables_pdf_ocr_and_keeps_layout_processing(monkeypatch):
    from docling.datamodel.base_models import ConversionStatus, InputFormat
    from docling.document_converter import DocumentConverter
    import pymupdf

    artifacts_path = r"C:\docling-artifacts"
    monkeypatch.setenv("DOCLING_ARTIFACTS_PATH", artifacts_path)
    monkeypatch.delenv("DOCLING_OCR_LANGUAGE", raising=False)
    pdf = pymupdf.open()
    pdf.new_page()
    pdf_bytes = pdf.tobytes()
    pdf.close()
    captured = {}

    def fake_init(self, **kwargs):
        captured["converter"] = self
        captured.update(kwargs)

    def fake_convert(self, source, *, raises_on_error):
        captured["source"] = source
        captured["raises_on_error"] = raises_on_error
        return SimpleNamespace(
            status=ConversionStatus.SUCCESS,
            document=SimpleNamespace(
                export_to_markdown=lambda: "Selectable PDF text",
                pages={},
            ),
        )

    monkeypatch.setattr(DocumentConverter, "__init__", fake_init)
    monkeypatch.setattr(DocumentConverter, "convert", fake_convert)

    parsed = DoclingParsingProvider().parse(
        RawSource(
            source_type="file",
            source_identifier="scanned.pdf",
            content_bytes=pdf_bytes,
        )
    )

    options = captured["format_options"][InputFormat.PDF].pipeline_options
    assert options.do_ocr is True
    assert options.ocr_options.backend == "onnxruntime"
    assert options.ocr_options.lang == ["en"]
    assert options.artifacts_path == artifacts_path
    assert captured["raises_on_error"] is False
    assert parsed.content == "Selectable PDF text"


def test_parser_provider_setting_is_normalized_and_rejected_when_unknown(
    monkeypatch,
):
    monkeypatch.setenv("DOCUMENT_PARSER_PROVIDER", "  DOCLING ")
    assert load_settings().providers.parser == "docling"

    monkeypatch.setenv("DOCUMENT_PARSER_PROVIDER", "unknown")
    with pytest.raises(ValueError, match="Unsupported DOCUMENT_PARSER_PROVIDER"):
        load_settings()


def _pipeline_for_provider(source_type: str, parser_provider: str):
    settings = SimpleNamespace(
        providers=SimpleNamespace(embeddings="nomic", parser=parser_provider),
        storage=SimpleNamespace(),
        openrouter=SimpleNamespace(),
        vector_store_table_name="chunks",
        vector_store_dimension=3,
    )
    request = SimpleNamespace(
        app=SimpleNamespace(state=SimpleNamespace(settings=settings))
    )
    return get_knowledge_ingestion_pipeline(
        source_type=source_type,
        request=request,
        session=object(),
    )


def test_python_docx_parser_extracts_paragraphs_and_table_rows():
    document = Document()
    document.add_paragraph("Clinic project overview")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Service"
    table.cell(0, 1).text = "Hours"
    buffer = BytesIO()
    document.save(buffer)

    parsed = PythonDocxParsingProvider().parse(
        RawSource(
            source_type="file",
            source_identifier="clinic.docx",
            content_bytes=buffer.getvalue(),
        )
    )

    assert "Clinic project overview" in parsed.content
    assert "Service | Hours" in parsed.content
    assert parsed.metadata["parser"] == "python-docx"


def test_docling_extracts_docx_headings_and_tables_into_sections():
    document = Document()
    document.add_heading("Clinic Guide", level=1)
    document.add_paragraph("Clinic hours are Monday to Friday.")
    document.add_heading("Services", level=2)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Service"
    table.cell(0, 1).text = "Hours"
    buffer = BytesIO()
    document.save(buffer)

    parsed = DoclingParsingProvider().parse(
        RawSource(
            source_type="file",
            source_identifier="clinic.docx",
            content_bytes=buffer.getvalue(),
        )
    )

    assert parsed.title == "Clinic Guide"
    assert parsed.metadata["parser"] == "docling"
    assert any(
        "Clinic Guide" in section.heading
        and "Monday to Friday" in section.content
        for section in parsed.structured_sections
    )
    assert any(
        "Services" in section.heading and "Service" in section.content
        for section in parsed.structured_sections
    )


def test_docling_conversion_failure_is_reported_without_native_fallback():
    with pytest.raises(ApplicationError) as exc_info:
        DoclingParsingProvider().parse(
            RawSource(
                source_type="file",
                source_identifier="invalid.docx",
                content_bytes=b"not a docx document",
            )
        )

    assert exc_info.value.code == "docling_conversion_failed"


def test_plain_text_formats_keep_text_parser():
    for source_type in ("txt", "md", "json"):
        pipeline = _pipeline_for(source_type)

        assert isinstance(pipeline.source_loader, FileSourceLoader)
        assert isinstance(pipeline.parser, TextDocumentParser)


def test_csv_keeps_text_source_loader():
    pipeline = _pipeline_for("csv")

    assert pipeline.source_loader.__class__.__name__ == "CsvSourceLoader"
    assert isinstance(pipeline.parser, CsvDocumentParser)


def test_pdf_uses_pymupdf_directly_without_fallback():
    pipeline = _pipeline_for("pdf")

    assert isinstance(pipeline.source_loader, FileSourceLoader)
    assert isinstance(pipeline.parser, StructuredDocumentParser)
    assert isinstance(pipeline.parser.parsing_contract, PyMuPDFParsingProvider)


def test_pymupdf_extracts_text_from_pdf_bytes():
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Knowledge base PDF fixture")
    pdf_bytes = pdf.tobytes()
    pdf.close()

    parsed = PyMuPDFParsingProvider().parse(
        RawSource(
            source_type="file",
            source_identifier="guide.pdf",
            content_bytes=pdf_bytes,
        )
    )

    assert "Knowledge base PDF fixture" in parsed.content
    assert parsed.metadata["parser"] == "pymupdf"


def test_pymupdf_extracts_selectable_text_and_ignores_image_only_pages():
    image_pdf = pymupdf.open()
    image_page = image_pdf.new_page()
    image_page.insert_text((72, 72), "Text that exists only in an image")
    image_pixmap = image_page.get_pixmap()
    image_pdf.close()

    pdf = pymupdf.open()
    mixed_page = pdf.new_page()
    mixed_page.insert_text((72, 72), "Selectable text remains indexable")
    mixed_page.insert_image(
        pymupdf.Rect(72, 100, 300, 300),
        pixmap=image_pixmap,
    )
    image_only_page = pdf.new_page()
    image_only_page.insert_image(
        pymupdf.Rect(72, 72, 300, 300),
        pixmap=image_pixmap,
    )
    pdf_bytes = pdf.tobytes()
    pdf.close()

    parsed = PyMuPDFParsingProvider().parse(
        RawSource(
            source_type="file",
            source_identifier="mixed-content.pdf",
            content_bytes=pdf_bytes,
        )
    )

    assert "Selectable text remains indexable" in parsed.content
    assert "Text that exists only in an image" not in parsed.content
    assert parsed.metadata["page_count"] == "2"
    assert parsed.metadata["pages_with_selectable_text"] == "1"


def test_native_pdf_runs_through_chunks_embedding_and_index(monkeypatch):
    from app.knowledge_engine.shared.models import EmbeddedChunk

    image_source_pdf = pymupdf.open()
    image_source_page = image_source_pdf.new_page()
    image_source_page.insert_text(
        (72, 72),
        "IMAGE_ONLY_WORDS_MUST_NOT_BE_OCRD",
    )
    image_pixmap = image_source_page.get_pixmap()
    image_source_pdf.close()

    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_textbox(
        pymupdf.Rect(50, 50, 550, 750),
        "\n\n".join(
            f"Support policy section {index}: "
            + ("Customers can contact support for account assistance. " * 8)
            for index in range(1, 8)
        ),
    )
    page.insert_image(
        pymupdf.Rect(50, 500, 150, 600),
        pixmap=image_pixmap,
    )
    image_only_page = pdf.new_page()
    image_only_page.insert_image(
        pymupdf.Rect(50, 50, 250, 250),
        pixmap=image_pixmap,
    )
    pdf_bytes = pdf.tobytes()
    pdf.close()
    source = RawSource(
        source_type="file",
        source_identifier="support.pdf",
        content_bytes=pdf_bytes,
    )
    pipeline = _pipeline_for_provider("pdf", "native")
    monkeypatch.setattr(
        pipeline.source_loader,
        "load",
        lambda _request: source,
    )

    class FakeEmbeddings(EmbeddingsContract):
        def embed_query(self, text):
            return [0.1, 0.2, 0.3]

        def embed_documents(self, texts):
            assert texts
            return [[0.1, 0.2, 0.3] for _ in texts]

    class FakeVectorStore:
        def __init__(self):
            self.indexed = []

        def index_chunks(self, chunks):
            self.indexed.extend(chunks)

        def verify_document_version(
            self,
            *,
            document_id,
            ingestion_version,
            expected_chunk_ids,
        ):
            indexed_ids = {
                chunk.chunk_id
                for chunk in self.indexed
                if chunk.metadata["document_id"] == document_id
                and int(chunk.metadata["ingestion_version"]) == ingestion_version
            }
            return indexed_ids == set(expected_chunk_ids)

    vector_store = FakeVectorStore()
    pipeline.embedding_generator.embeddings_contract = FakeEmbeddings()
    pipeline.vector_indexer.vector_store_contract = cast(
        VectorStoreContract,
        vector_store,
    )

    result = pipeline.run(
        KnowledgeIngestionPipelineRequest(
            document_id="native-pdf-document",
            knowledge_base_id="knowledge-base",
            source_type="pdf",
            source_path="support.pdf",
            source_identifier="support.pdf",
            ingestion_version=1,
        )
    )

    assert result.chunk_count > 0
    assert len(vector_store.indexed) == result.chunk_count
    assert all(isinstance(chunk, EmbeddedChunk) for chunk in vector_store.indexed)
    assert all(chunk.embedding == [0.1, 0.2, 0.3] for chunk in vector_store.indexed)
    assert "Support policy section" in " ".join(
        chunk.content for chunk in vector_store.indexed
    )
    assert "IMAGE_ONLY_WORDS_MUST_NOT_BE_OCRD" not in " ".join(
        chunk.content for chunk in vector_store.indexed
    )


def test_scanned_pdf_without_selectable_text_fails_with_clear_error(caplog):
    pdf = pymupdf.open()
    source_pdf = pymupdf.open()
    source_page = source_pdf.new_page()
    source_page.insert_text((72, 72), "Text available only as image pixels")
    image_pixmap = source_page.get_pixmap()
    source_pdf.close()
    image_page = pdf.new_page()
    image_page.insert_image(
        pymupdf.Rect(72, 72, 500, 700),
        pixmap=image_pixmap,
    )
    pdf_bytes = pdf.tobytes()
    pdf.close()

    with pytest.raises(ApplicationError) as exc_info:
        PyMuPDFParsingProvider().parse(
            RawSource(
                source_type="file",
                source_identifier="scanned-guide.pdf",
                content_bytes=pdf_bytes,
            )
        )

    assert exc_info.value.code == "pdf_text_not_found"
    assert "scanned" in exc_info.value.message.lower()
    assert (
        "pages=1, pages_with_selectable_text=0, "
        "extracted_characters=0"
    ) in caplog.text


@pytest.mark.parametrize(
    "source_type",
    ["txt", "md", "json"],
)
def test_text_upload_formats_extract_content(source_type):
    parsed = TextDocumentParser().parse(
        RawSource(
            source_type=source_type,
            source_identifier=f"guide.{source_type}",
            content_bytes=b'{"answer": "The office opens at 9."}',
        )
    )

    assert "office opens at 9" in parsed.content
    assert parsed.title == f"guide.{source_type}"


@pytest.mark.parametrize(
    ("extension", "expected_type"),
    [
        ("pdf", "pdf"),
        ("docx", "docx"),
        ("txt", "txt"),
        ("csv", "csv"),
        ("json", "json"),
        ("md", "md"),
    ],
)
def test_uploaded_file_type_is_resolved_from_storage_path(extension, expected_type):
    document = SimpleNamespace(source_type="file")

    assert (
        _resolve_source_type(
            document,
            f"kb/document/source.{extension}",
        )
        == expected_type
    )


def test_unsupported_office_format_fails_fast():
    with pytest.raises(ValueError, match="Unsupported ingestion source type: xlsx"):
        _pipeline_for("xlsx")


def test_ingestion_does_not_publish_documents_without_chunks(monkeypatch):
    def fail_if_called(*_args, **_kwargs):
        pytest.fail("Ingestion continued after chunk generation returned no chunks.")

    pipeline = _pipeline_for_provider("pdf", "native")
    monkeypatch.setattr(
        pipeline.source_loader,
        "load",
        lambda _request: RawSource(
            source_type="file",
            source_identifier="empty.pdf",
            content_bytes=b"source",
        ),
    )
    monkeypatch.setattr(
        pipeline.parser,
        "parse",
        lambda _source: ParsedDocument(
            title="Empty extraction",
            content="some extracted text",
            sections=["some extracted text"],
        ),
    )
    monkeypatch.setattr(
        pipeline.chunk_generator,
        "generate",
        lambda **_kwargs: [],
    )
    monkeypatch.setattr(
        type(pipeline.metadata_enricher),
        "enrich",
        fail_if_called,
    )
    monkeypatch.setattr(
        type(pipeline.embedding_generator),
        "generate",
        fail_if_called,
    )
    monkeypatch.setattr(
        type(pipeline.vector_indexer),
        "index",
        fail_if_called,
    )

    request = KnowledgeIngestionPipelineRequest(
        document_id="document-id",
        knowledge_base_id="knowledge-base-id",
        source_type="pdf",
        source_path="empty.pdf",
        source_identifier="empty.pdf",
        ingestion_version=1,
    )

    with pytest.raises(ApplicationError) as exc_info:
        pipeline.run(request)

    assert exc_info.value.code == "document_chunks_empty"


def test_csv_rows_keep_column_names_with_values():
    parsed = CsvDocumentParser().parse(
        RawSource(
            source_type="csv",
            source_identifier="support.csv",
            content_text="plan,price\nBasic,10\nPro,25",
        )
    )

    assert "plan: Basic" in parsed.content
    assert "price: 25" in parsed.content


def test_csv_parser_handles_utf8_bom():
    parsed = CsvDocumentParser().parse(
        RawSource(
            source_type="csv",
            source_identifier="support.csv",
            content_bytes="\ufeffplan,price\nBasic,10".encode("utf-8"),
        )
    )

    assert "plan: Basic" in parsed.content


def test_upload_api_supports_the_formats_shown_in_the_registry():
    assert SUPPORTED_UPLOAD_EXTENSIONS == {
        ".pdf",
        ".docx",
        ".txt",
        ".csv",
        ".json",
        ".md",
    }
