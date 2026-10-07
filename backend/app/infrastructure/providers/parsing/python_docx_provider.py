from __future__ import annotations

import logging
import time
from dataclasses import dataclass
from io import BytesIO

from app.core.exceptions import ApplicationError
from app.knowledge_engine.contracts.parsing import ParsingContract
from app.knowledge_engine.shared.helpers import (
    normalize_whitespace,
    split_text_into_paragraphs,
)
from app.knowledge_engine.shared.models import ParsedDocument, RawSource

logger = logging.getLogger(__name__)


@dataclass(slots=True)
class PythonDocxParsingProvider(ParsingContract):
    def parse(self, source: RawSource) -> ParsedDocument:
        if source.content_bytes is None:
            raise ApplicationError(
                message="DOCX binary content is missing.",
                code="docx_content_missing",
                status_code=422,
            )

        try:
            from docx import Document

            started_at = time.monotonic()
            logger.info("DOCX parser opening package")
            document = Document(BytesIO(source.content_bytes))
            logger.info(
                "DOCX package opened: paragraphs=%s tables=%s duration_seconds=%.3f",
                len(document.paragraphs),
                len(document.tables),
                time.monotonic() - started_at,
            )
        except Exception as exc:
            raise ApplicationError(
                message=f"DOCX parsing failed: {exc}",
                code="docx_parsing_failed",
                status_code=422,
            ) from exc

        content_blocks = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]
        logger.info("DOCX paragraphs extracted: count=%s", len(content_blocks))
        table_started_at = time.monotonic()
        extracted_table_rows = 0
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    content_blocks.append(" | ".join(cells))
                    extracted_table_rows += 1
        logger.info(
            "DOCX table rows extracted: count=%s duration_seconds=%.3f",
            extracted_table_rows,
            time.monotonic() - table_started_at,
        )

        content = normalize_whitespace("\n\n".join(content_blocks))
        if not content:
            raise ApplicationError(
                message="No text was found in the DOCX document.",
                code="docx_content_empty",
                status_code=422,
            )

        source_name = source.source_identifier.rsplit("/", 1)[-1] or "document.docx"
        title = source.metadata.get("document_title") or source_name
        paragraphs = split_text_into_paragraphs(content)

        return ParsedDocument(
            title=title,
            content=content,
            sections=paragraphs or [content],
            metadata={**source.metadata, "parser": "python-docx"},
        )